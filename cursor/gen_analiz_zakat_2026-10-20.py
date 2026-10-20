#!/usr/bin/env python3
"""Анализ РК «Закат на море» 20.10.2026 — много показов, низкий CTR → Word."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Analiz_RK_Zakat_na_more_20_2026-10-20.docx"


def set_run_font(run, size=11, bold=False, color=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(
            run,
            size=16 if level == 1 else (13 if level == 2 else 12),
            bold=True,
            color=RGBColor(0x0B, 0x3D, 0x5C),
        )
    return p


def add_p(doc, text, *, bold=False, size=11, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(2)


def shade_cell(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade_cell(cell, "0B3D5C")
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
            if r_i % 2 == 1:
                shade_cell(cell, "F2F7FA")
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def add_minus_lines(doc, phrases):
    for phrase in phrases:
        p = doc.add_paragraph()
        run = p.add_run(phrase)
        set_run_font(run, size=11)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(1.8)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("АНАЛИЗ КАМПАНИИ")
    set_run_font(r, size=12, bold=True, color=RGBColor(0x0B, 0x3D, 0x5C))

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Закат на море — 20.10, много показов")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x0B, 0x3D, 0x5C))

    for line in [
        "Группа 5779888441 · vitaminki21 · код ZK · только 20.10 (до 20:31)",
        "CSV 2026-10-20_20-31-44_vitaminki21 · Поиск + РСЯ",
        "Стратегию, бюджет и цели не менять",
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        set_run_font(r, size=10, color=RGBColor(0x44, 0x44, 0x44))

    add_heading(doc, "1. Короткий вердикт", 1)
    add_bullets(
        doc,
        [
            "20.10: 0 ₽ · 285 показов · 5 кликов · 0 конв. · CTR 1,75%. Это один день — больше, чем вся суббота 17.10 (176).",
            "99% объёма — РСЯ (282/5, CTR 1,77%). Поиск 3/0. Автотаргет 100%. Ключи-фразы снова 0 показов.",
            "Низкий CTR — не «сломались объявления». Сети льют дешёвые показы «Прочее»: десктоп 103/0, 55+ 103/1, мужчины 221/3.",
            "Рабочий кусок: смартфоны 170/4 CTR 2,35%; заголовок «Вечерняя прогулка, яхта только для вас» 33/2 CTR 6,1%. Его мало крутят.",
            "Деньги не горят (0 ₽). Стратегию не трогать. РСЯ не выключать. Резать корректировками десктоп и 55+.",
        ],
    )

    add_heading(doc, "2. Почему столько показов", 1)
    add_p(
        doc,
        "После включения РСЯ 28.08 Директ набирает дешёвый охват в сетях. "
        "Оплата за конверсии: показов много, расход 0, пока нет ценности. "
        "17.10 уже был скачок 176 показов. 20.10 — 285 за день. Это не всплеск спроса на закат, это инвентарь РСЯ.",
    )
    add_table(
        doc,
        ["Срез", "Показы", "Клики", "CTR", "Расход", "Конв."],
        [
            ["05–10.10 (6 дн.)", "159", "6", "3,77%", "156 ₽", "1"],
            ["12–17.10 (6 дн.)", "381", "6", "1,57%", "0 ₽", "0"],
            ["17.10 сб", "176", "2", "1,1%", "0 ₽", "0"],
            ["20.10 вт до 20:31", "285", "5", "1,75%", "0 ₽", "0"],
        ],
        col_widths=[4.5, 2.4, 2.2, 2.4, 2.4, 2.2],
    )

    add_heading(doc, "3. Где просел CTR", 1)
    add_table(
        doc,
        ["Срез", "Показы", "Клики", "CTR"],
        [
            ["РСЯ", "282", "5", "1,77%"],
            ["Поиск", "3", "0", "0%"],
            ["Смартфоны", "170", "4", "2,35%"],
            ["Десктопы", "103", "0", "0%"],
            ["Мужчины", "221", "3", "1,36%"],
            ["Женщины", "63", "2", "3,17%"],
            ["25–34", "39", "2", "5,13%"],
            ["35–44", "84", "0", "0%"],
            ["45–54", "49", "2", "4,08%"],
            ["Старше 55", "103", "1", "0,97%"],
            ["Остальные (платёж.)", "271", "5", "1,85%"],
        ],
        col_widths=[5.5, 3, 3, 3],
    )
    add_p(
        doc,
        "Десктоп = 36% показов и 0 кликов. 55+ = 36% показов и 1 клик. "
        "ЦА 25–44 на смартфоне живее. 35–44 обнулили десктопные мужчины (28/0 десктоп + 35/0 смарт муж 35–44).",
        bold=True,
    )

    add_heading(doc, "4. Креативы: объём не туда", 1)
    add_table(
        doc,
        ["Заголовок", "Показы", "Клики", "CTR"],
        [
            ["Бронируйте… всю яхту — с 18:00", "108", "2", "1,85%"],
            ["Бронируйте… всю яхту — с 17:00", "49", "1", "2,04%"],
            ["Выход с 17:00 часов. Сириус", "38", "0", "0%"],
            ["Вечерняя прогулка, яхта только для вас", "33", "2", "6,06%"],
            ["Катеров мало", "21", "0", "0%"],
        ],
        col_widths=[8.5, 2.5, 2.2, 2.2],
    )
    add_table(
        doc,
        ["Картинка", "Показы", "Клики", "CTR"],
        [
            ["закат 05.jpg", "111", "3", "2,70%"],
            ["закат 02.jpg", "101", "2", "1,98%"],
            ["закат в море.jpg", "47", "0", "0%"],
            ["закат 2.jpg (конв. 09.10)", "14", "0", "0%"],
        ],
        col_widths=[8.5, 2.5, 2.2, 2.2],
    )
    add_p(
        doc,
        "Комбинатор крутит «всю яхту с 18:00» на десктоп (39/0). "
        "Заголовок с CTR 6% получает втрое меньше показов. Тексты не переписывать — крутить «Вечерняя прогулка».",
    )

    add_heading(doc, "5. Поиск — 3 показа", 1)
    add_table(
        doc,
        ["Запрос", "П/К", "Вердикт"],
        [
            ["закат на яхте в акватории черного моря", "1/0", "Ядро. Ключ"],
            ["аренда большой яхты сириус", "1/0", "SR, не закат"],
            ["путешествие на яхте в россии", "1/0", "Минус"],
        ],
        col_widths=[8.5, 2.2, 5.5],
    )
    add_p(
        doc,
        "Фразы 10.10 снова не вышли. Единственный закатный запрос дня поймал автотаргет, не ключ.",
        bold=True,
    )

    add_heading(doc, "6. Что сделать", 1)
    add_heading(doc, "6.1. Не трогать", 2)
    add_bullets(
        doc,
        [
            "Стратегия, бюджет 12 000, цели 140/130/110/120 — нет.",
            "РСЯ не выключать: 09.10 она дала конверсию, расход сегодня 0.",
            "Галерею и профиль организации не включать.",
        ],
    )
    add_heading(doc, "6.2. Корректировки (CTR, не стратегия)", 2)
    add_bullets(
        doc,
        [
            "Десктоп −50…−100% (103 показа, 0 кликов).",
            "Старше 55 −50% (103/1).",
            "Не резать смартфоны 25–54.",
        ],
    )
    add_heading(doc, "6.3. Минус", 2)
    add_p(doc, "Столбиком:")
    add_minus_lines(doc, ["путешествие", "катание"])
    add_p(doc, "Не минусовать: закат, вечерняя, для двоих, целиком, аренда, ужин, морпорт.")
    add_heading(doc, "6.4. Ключ", 2)
    add_bullets(doc, ['"закат на яхте в акватории черного моря"'])
    add_p(doc, "Пачку 10.10 не дублировать — сначала проверить, что она в группе 5779888441 и не на паузе.")

    add_heading(doc, "7. Чего не делать", 1)
    add_bullets(
        doc,
        [
            "Не поднимать цели «потому что CTR низкий» — расход 0, конверсий 0.",
            "Не выключать РСЯ из‑за пачки показов.",
            "Не переписывать объявления: проблема в ротации и аудитории сетей, не в тексте 17:00–19:00.",
            "Не вести на /progulki_na_yacht, /delfin, не писать 1 800 / 2 500.",
        ],
    )

    add_heading(doc, "8. Источники", 1)
    add_bullets(
        doc,
        [
            "CSV: 2026-10-20_20-31-44_vitaminki21.csv (группа 5779888441, только 20.10)",
            "Сравнение: 12–17.10 и 05–10.10",
            "Память: cursor/Zakat_campaign_memory.md",
        ],
    )

    doc.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    build()
